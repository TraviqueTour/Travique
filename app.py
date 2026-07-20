"""
Travique — Discover Chiang Mai
A simple local tour-booking website (no heavy web framework required).

What it does:
  * Shows every tour from the booklet as a full-screen "slide".
  * Guests pick a date + time slot and book.
  * Each booking is appended to bookings.xlsx (openable in Excel).
  * A confirmation email is sent to the guest (via Resend, if configured).
  * /admin shows all bookings at a glance.

Run:  python3 app.py   then open http://127.0.0.1:5000
"""
import os
import re
import json
import datetime as dt
import mimetypes
from threading import Lock
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

from jinja2 import Template
from openpyxl import Workbook, load_workbook

from tours import TOURS, TOURS_BY_ID
from templates import INDEX_HTML, ADMIN_HTML, EMAIL_HTML

# ---------------------------------------------------------------- config
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
XLSX_PATH = os.path.join(BASE_DIR, "bookings.xlsx")
HOST, PORT = "127.0.0.1", 5000
_lock = Lock()

# Email (optional). Set as environment variables to enable real emails.
RESEND_API_KEY = os.environ.get("RESEND_API_KEY", "").strip()
FROM_EMAIL = os.environ.get("FROM_EMAIL", "Travique <onboarding@resend.dev>").strip()
NOTIFY_EMAIL = os.environ.get("NOTIFY_EMAIL", "").strip()

HEADERS = [
    "Booking ID", "Booked At", "Tour", "Tour Date", "Time Slot",
    "Guest Name", "Email", "Phone", "Adults", "Children",
    "Estimated Total (THB)", "Notes", "Email Status",
]

_index_t = Template(INDEX_HTML)
_admin_t = Template(ADMIN_HTML)
_email_t = Template(EMAIL_HTML)


# ---------------------------------------------------------------- excel
def _ensure_workbook():
    if not os.path.exists(XLSX_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Bookings"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = c.font.copy(bold=True)
        widths = [16, 20, 28, 13, 26, 22, 28, 16, 8, 9, 20, 30, 16]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        wb.save(XLSX_PATH)


def append_booking(row):
    with _lock:
        _ensure_workbook()
        wb = load_workbook(XLSX_PATH)
        ws = wb["Bookings"]
        ws.append([row.get(h, "") for h in HEADERS])
        wb.save(XLSX_PATH)


def read_bookings():
    if not os.path.exists(XLSX_PATH):
        return []
    wb = load_workbook(XLSX_PATH, read_only=True)
    ws = wb["Bookings"]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(HEADERS, r)) for r in rows[1:]]


# ---------------------------------------------------------------- email
def send_confirmation(to_email, guest_name, tour, date_str, slot, total):
    if not RESEND_API_KEY:
        return "logged (no email key)"
    html = _email_t.render(guest=guest_name, tour=tour, date_str=date_str,
                           slot=slot, total=total)
    subject = f"Your Travique booking - {tour['title']}"
    try:
        import requests
        recipients = [to_email] + ([NOTIFY_EMAIL] if NOTIFY_EMAIL else [])
        resp = requests.post(
            "https://api.resend.com/emails",
            headers={"Authorization": f"Bearer {RESEND_API_KEY}",
                     "Content-Type": "application/json"},
            json={"from": FROM_EMAIL, "to": recipients,
                  "subject": subject, "html": html},
            timeout=15,
        )
        return "sent" if resp.status_code in (200, 201) else f"failed ({resp.status_code})"
    except Exception as e:  # noqa
        return f"error ({type(e).__name__})"


# ---------------------------------------------------------------- booking logic
def process_booking(d):
    """Validate + record a booking. Returns (status_code, response_dict)."""
    tour = TOURS_BY_ID.get(d.get("tour_id"))
    if not tour:
        return 400, {"ok": False, "error": "Unknown tour."}

    name = (d.get("name") or "").strip()
    email = (d.get("email") or "").strip()
    if not name or not re.match(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return 400, {"ok": False, "error": "Please enter a name and a valid email."}

    date_str = (d.get("date") or "").strip()
    if not date_str:
        return 400, {"ok": False, "error": "Please choose a tour date."}
    slot = (d.get("slot") or "").strip()
    phone = (d.get("phone") or "").strip()
    notes = (d.get("notes") or "").strip()

    def _int(v, lo, default):
        try:
            return max(lo, int(v))
        except (TypeError, ValueError):
            return default
    adults = _int(d.get("adults"), 1, 1)
    children = _int(d.get("children"), 0, 0)

    kid_price = tour["kid_price"] if tour["kid_price"] is not None else tour["price"]
    total = adults * tour["price"] + children * kid_price

    booking_id = "TRV-" + dt.datetime.now().strftime("%y%m%d-%H%M%S-") + \
        dt.datetime.now().strftime("%f")[:3]
    email_status = send_confirmation(email, name, tour, date_str, slot, total)

    append_booking({
        "Booking ID": booking_id,
        "Booked At": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Tour": tour["title"], "Tour Date": date_str, "Time Slot": slot,
        "Guest Name": name, "Email": email, "Phone": phone,
        "Adults": adults, "Children": children,
        "Estimated Total (THB)": total, "Notes": notes,
        "Email Status": email_status,
    })
    return 200, {"ok": True, "booking_id": booking_id, "total": total,
                 "email_status": email_status}


def render_index():
    return _index_t.render(tours=TOURS, tours_json=json.dumps(TOURS),
                           today=dt.date.today().isoformat())


def render_admin():
    bookings = list(reversed(read_bookings()))
    total_guests = sum((int(b.get("Adults") or 0) + int(b.get("Children") or 0))
                       for b in bookings)
    revenue = sum(int(b.get("Estimated Total (THB)") or 0) for b in bookings)
    return _admin_t.render(bookings=bookings, total_guests=total_guests,
                           revenue=revenue, xlsx=os.path.basename(XLSX_PATH))


# ---------------------------------------------------------------- http server
class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):  # quieter console
        pass

    def _send(self, code, body, ctype="text/html; charset=utf-8"):
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/":
            return self._send(200, render_index())
        if path == "/admin":
            return self._send(200, render_admin())
        if path.startswith("/static/"):
            return self._serve_static(path)
        return self._send(404, "<h1>404 Not Found</h1>")

    def _serve_static(self, path):
        rel = path[len("/static/"):]
        full = os.path.normpath(os.path.join(STATIC_DIR, rel))
        if not full.startswith(STATIC_DIR) or not os.path.isfile(full):
            return self._send(404, "not found", "text/plain")
        ctype = mimetypes.guess_type(full)[0] or "application/octet-stream"
        with open(full, "rb") as f:
            data = f.read()
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "public, max-age=86400")
        self.end_headers()
        self.wfile.write(data)

    def do_POST(self):
        if urlparse(self.path).path != "/book":
            return self._send(404, "not found", "text/plain")
        try:
            length = int(self.headers.get("Content-Length", 0))
            data = json.loads(self.rfile.read(length) or b"{}")
        except Exception:
            return self._send(400, json.dumps({"ok": False, "error": "Bad request."}),
                              "application/json")
        code, resp = process_booking(data)
        self._send(code, json.dumps(resp), "application/json")


def main():
    _ensure_workbook()
    print("\n  Travique is running ->  http://127.0.0.1:%d" % PORT)
    print("  Bookings view       ->  http://127.0.0.1:%d/admin" % PORT)
    print("  Excel file          ->  " + XLSX_PATH)
    print("  Emails: " + ("ON (Resend)" if RESEND_API_KEY else "logging only (no key set)"))
    print("\n  Press Ctrl+C to stop.\n")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()


if __name__ == "__main__":
    main()
